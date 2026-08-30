from io import BytesIO
from uuid import UUID

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, StreamObject
from sqlalchemy import select
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.ext.asyncio import AsyncSession

from app.common.config import Settings
from app.common.db_models import DocumentChunk, File, FileProcessingStatus
from app.common.errors import ApiError
from app.modules.files.file_repository import FileRepository
from app.modules.files.file_service import FileService
from tests.conftest import DEV_SUPERVISOR_ID


class FakeStorage:
    def __init__(self, events: list[str] | None = None) -> None:
        self.uploaded_paths: list[str] = []
        self.deleted_paths: list[str] = []
        self.events = events

    def upload_bytes(
        self,
        storage_path: str,
        content: bytes,
        content_type: str,
        metadata: dict[str, str] | None = None,
    ) -> None:
        self.uploaded_paths.append(storage_path)

    def delete_object(self, storage_path: str) -> None:
        if self.events is not None:
            self.events.append("storage_delete")
        self.deleted_paths.append(storage_path)

    def download_bytes(self, storage_path: str) -> bytes:
        return b"stored text"

    def presigned_download_url(self, storage_path: str, expires_in_seconds: int = 3600) -> str:
        return f"http://storage.local/{storage_path}?expires={expires_in_seconds}"


class FailingUploadStorage(FakeStorage):
    def upload_bytes(
        self,
        storage_path: str,
        content: bytes,
        content_type: str,
        metadata: dict[str, str] | None = None,
    ) -> None:
        msg = "storage upload failed"
        raise RuntimeError(msg)


class FailingDeleteStorage(FakeStorage):
    def delete_object(self, storage_path: str) -> None:
        msg = "storage delete failed"
        raise RuntimeError(msg)


class FailingDownloadStorage(FakeStorage):
    def presigned_download_url(self, storage_path: str, expires_in_seconds: int = 3600) -> str:
        msg = "storage download URL failed"
        raise RuntimeError(msg)


class FakeEmbeddings:
    async def embed_documents(self, texts: list[str]) -> list[list[float]]:
        return [[float(index)] * 1024 for index, _text in enumerate(texts)]

    async def embed_query(self, text: str) -> list[float]:
        return [0.1] * 1024


class WrongDimensionEmbeddings:
    async def embed_documents(self, texts: list[str]) -> list[list[float]]:
        return [[0.1, 0.2, 0.3] for _text in texts]

    async def embed_query(self, text: str) -> list[float]:
        return [0.1, 0.2, 0.3]


class FakeImageDescriptions:
    def __init__(self, description: str = "A labelled architecture diagram.") -> None:
        self.description = description
        self.requests: list[tuple[bytes, str]] = []

    async def describe_image(self, content: bytes, mime_type: str) -> str:
        self.requests.append((content, mime_type))
        return self.description


class FailingImageDescriptions(FakeImageDescriptions):
    async def describe_image(self, content: bytes, mime_type: str) -> str:
        raise RuntimeError("vision model failed")


def _pdf_with_text(text: str) -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): writer._add_object(font)}  # noqa: SLF001
            )
        }
    )
    stream = StreamObject()
    stream.set_data(f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode())
    page[NameObject("/Contents")] = writer._add_object(stream)  # noqa: SLF001
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


class FailingCreateRepository(FileRepository):
    async def create(self, **_kwargs):
        msg = "database write failed"
        raise SQLAlchemyError(msg)


class DuplicateChunkRepository(FileRepository):
    async def add_chunks(
        self, file_id: UUID, chunks: list[str], embeddings: list[list[float]]
    ) -> None:
        for chunk in chunks:
            self.session.add(
                DocumentChunk(
                    file_id=file_id,
                    chunk_index=0,
                    content=chunk,
                    page_number=None,
                    embedding=embeddings[0],
                )
            )
        await self.session.flush()


class RecordingFileRepository(FileRepository):
    def __init__(self, session: AsyncSession, events: list[str]) -> None:
        super().__init__(session)
        self.events = events

    async def delete(self, file: File) -> None:
        await super().delete(file)
        self.events.append("repository_delete")


@pytest.mark.asyncio
async def test_upload_text_file_stores_file_and_chunks(session: AsyncSession) -> None:
    storage = FakeStorage()
    settings = Settings(chunk_size=20, chunk_overlap=5)
    service = FileService(session, storage=storage, embeddings=FakeEmbeddings(), settings=settings)

    result = await service.upload_file(
        supervisor_id=DEV_SUPERVISOR_ID,
        original_filename="notes.txt",
        content=b"Capstone teams need regular supervisor feedback.",
        mime_type="text/plain",
    )

    assert result.processing_status == FileProcessingStatus.READY
    assert result.file_size_bytes == len(b"Capstone teams need regular supervisor feedback.")
    assert len(storage.uploaded_paths) == 1

    chunks = list(
        await session.scalars(select(DocumentChunk).where(DocumentChunk.file_id == result.id))
    )
    assert len(chunks) > 1


@pytest.mark.asyncio
async def test_upload_rejects_unsupported_file(session: AsyncSession) -> None:
    service = FileService(
        session,
        storage=FakeStorage(),
        embeddings=FakeEmbeddings(),
        settings=Settings(),
    )

    with pytest.raises(ApiError) as exc:
        await service.upload_file(
            supervisor_id=UUID("11111111-1111-1111-1111-111111111111"),
            original_filename="archive.zip",
            content=b"not text",
            mime_type="application/zip",
        )

    assert exc.value.code == "UNSUPPORTED_FILE_TYPE"


@pytest.mark.asyncio
async def test_upload_marks_file_failed_when_embedding_dimension_is_wrong(
    session: AsyncSession,
) -> None:
    storage = FakeStorage()
    service = FileService(
        session,
        storage=storage,
        embeddings=WrongDimensionEmbeddings(),
        settings=Settings(embedding_dimension=1024),
    )

    with pytest.raises(ApiError) as exc:
        await service.upload_file(
            supervisor_id=DEV_SUPERVISOR_ID,
            original_filename="notes.txt",
            content=b"Valid text.",
            mime_type="text/plain",
        )

    assert exc.value.code == "DOCUMENT_PROCESSING_FAILED"
    assert exc.value.http_status == 502


@pytest.mark.asyncio
async def test_upload_marks_file_failed_when_chunk_database_write_fails(
    session: AsyncSession,
) -> None:
    storage = FakeStorage()
    service = FileService(
        session,
        repository=DuplicateChunkRepository(session),
        storage=storage,
        embeddings=FakeEmbeddings(),
        settings=Settings(chunk_size=10, chunk_overlap=0),
    )

    with pytest.raises(ApiError) as exc:
        await service.upload_file(
            supervisor_id=DEV_SUPERVISOR_ID,
            original_filename="notes.txt",
            content=b"First chunk. Second chunk.",
            mime_type="text/plain",
        )

    files = list(await session.scalars(select(File)))
    chunks = list(await session.scalars(select(DocumentChunk)))
    assert exc.value.code == "DOCUMENT_PROCESSING_FAILED"
    assert exc.value.http_status == 502
    assert len(files) == 1
    assert files[0].processing_status == FileProcessingStatus.FAILED
    assert chunks == []


@pytest.mark.asyncio
async def test_upload_rejects_oversized_content_before_storage_write(
    session: AsyncSession,
) -> None:
    storage = FakeStorage()
    service = FileService(
        session,
        storage=storage,
        embeddings=FakeEmbeddings(),
        settings=Settings(max_upload_bytes=4),
    )

    with pytest.raises(ApiError) as exc:
        await service.upload_file(
            supervisor_id=DEV_SUPERVISOR_ID,
            original_filename="notes.txt",
            content=b"too large",
            mime_type="text/plain",
        )

    assert exc.value.code == "FILE_TOO_LARGE"
    assert exc.value.http_status == 413
    assert storage.uploaded_paths == []


@pytest.mark.asyncio
async def test_upload_shapes_storage_upload_failure_before_database_write(
    session: AsyncSession,
) -> None:
    service = FileService(
        session,
        storage=FailingUploadStorage(),
        embeddings=FakeEmbeddings(),
        settings=Settings(),
    )

    with pytest.raises(ApiError) as exc:
        await service.upload_file(
            supervisor_id=DEV_SUPERVISOR_ID,
            original_filename="notes.txt",
            content=b"Valid text.",
            mime_type="text/plain",
        )

    files = list(await session.scalars(select(File)))
    assert exc.value.code == "OBJECT_STORAGE_UPLOAD_FAILED"
    assert exc.value.http_status == 502
    assert files == []


@pytest.mark.asyncio
async def test_upload_deletes_object_when_database_create_fails(
    session: AsyncSession,
) -> None:
    storage = FakeStorage()
    service = FileService(
        session,
        repository=FailingCreateRepository(session),
        storage=storage,
        embeddings=FakeEmbeddings(),
        settings=Settings(),
    )

    with pytest.raises(ApiError) as exc:
        await service.upload_file(
            supervisor_id=DEV_SUPERVISOR_ID,
            original_filename="notes.txt",
            content=b"Valid text.",
            mime_type="text/plain",
        )

    assert exc.value.code == "DATABASE_WRITE_FAILED"
    assert exc.value.http_status == 500
    assert storage.uploaded_paths == storage.deleted_paths


@pytest.mark.asyncio
async def test_upload_shapes_cleanup_failure_after_database_create_fails(
    session: AsyncSession,
) -> None:
    service = FileService(
        session,
        repository=FailingCreateRepository(session),
        storage=FailingDeleteStorage(),
        embeddings=FakeEmbeddings(),
        settings=Settings(),
    )

    with pytest.raises(ApiError) as exc:
        await service.upload_file(
            supervisor_id=DEV_SUPERVISOR_ID,
            original_filename="notes.txt",
            content=b"Valid text.",
            mime_type="text/plain",
        )

    assert exc.value.code == "OBJECT_STORAGE_CLEANUP_FAILED"
    assert exc.value.http_status == 502


@pytest.mark.asyncio
async def test_delete_flushes_database_before_storage_delete(session: AsyncSession) -> None:
    file = File(
        supervisor_id=DEV_SUPERVISOR_ID,
        original_filename="notes.txt",
        stored_filename="notes.txt",
        storage_base_url="http://localhost:9000/capstone-files/",
        storage_path="supervisors/dev/files/notes.txt",
        mime_type="text/plain",
        file_size_bytes=64,
        processing_status=FileProcessingStatus.READY,
    )
    session.add(file)
    await session.commit()

    events: list[str] = []
    service = FileService(
        session,
        repository=RecordingFileRepository(session, events),
        storage=FakeStorage(events),
        embeddings=FakeEmbeddings(),
        settings=Settings(),
    )

    await service.delete_file(DEV_SUPERVISOR_ID, file.id)

    assert events == ["repository_delete", "storage_delete"]


@pytest.mark.asyncio
async def test_delete_rolls_back_database_when_storage_delete_fails(
    session: AsyncSession,
) -> None:
    file = File(
        supervisor_id=DEV_SUPERVISOR_ID,
        original_filename="notes.txt",
        stored_filename="notes.txt",
        storage_base_url="http://localhost:9000/capstone-files/",
        storage_path="supervisors/dev/files/notes.txt",
        mime_type="text/plain",
        file_size_bytes=64,
        processing_status=FileProcessingStatus.READY,
    )
    session.add(file)
    await session.commit()
    service = FileService(
        session,
        storage=FailingDeleteStorage(),
        embeddings=FakeEmbeddings(),
        settings=Settings(),
    )

    with pytest.raises(ApiError) as exc:
        await service.delete_file(DEV_SUPERVISOR_ID, file.id)

    await session.refresh(file)
    assert exc.value.code == "OBJECT_STORAGE_DELETE_FAILED"
    assert exc.value.http_status == 502
    assert file.deleted_at is None


@pytest.mark.asyncio
async def test_download_url_shapes_storage_failure(session: AsyncSession) -> None:
    file = File(
        supervisor_id=DEV_SUPERVISOR_ID,
        original_filename="notes.txt",
        stored_filename="notes.txt",
        storage_base_url="http://localhost:9000/capstone-files/",
        storage_path="supervisors/dev/files/notes.txt",
        mime_type="text/plain",
        file_size_bytes=64,
        processing_status=FileProcessingStatus.READY,
    )
    session.add(file)
    await session.commit()
    service = FileService(
        session,
        storage=FailingDownloadStorage(),
        embeddings=FakeEmbeddings(),
        settings=Settings(),
    )

    with pytest.raises(ApiError) as exc:
        await service.get_download_url(DEV_SUPERVISOR_ID, file.id)

    assert exc.value.code == "OBJECT_STORAGE_DOWNLOAD_URL_FAILED"
    assert exc.value.http_status == 502


@pytest.mark.asyncio
async def test_upload_text_file_caps_stored_filename_length(session: AsyncSession) -> None:
    service = FileService(
        session,
        storage=FakeStorage(),
        embeddings=FakeEmbeddings(),
        settings=Settings(),
    )
    long_filename = f"{'a' * 251}.txt"

    result = await service.upload_file(
        supervisor_id=DEV_SUPERVISOR_ID,
        original_filename=long_filename,
        content=b"Long filename content.",
        mime_type="text/plain",
    )

    file = await session.get(File, result.id)
    assert file is not None
    assert len(file.stored_filename) <= 255
    assert file.stored_filename.endswith(".txt")


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("original_filename", "mime_type", "field"),
    [
        (f"{'a' * 256}.txt", "text/plain", "file.filename"),
        ("notes.txt", f"text/{'a' * 251}", "file.content_type"),
    ],
)
async def test_upload_rejects_overlong_metadata_before_storage_write(
    session: AsyncSession,
    original_filename: str,
    mime_type: str,
    field: str,
) -> None:
    storage = FakeStorage()
    service = FileService(
        session,
        storage=storage,
        embeddings=FakeEmbeddings(),
        settings=Settings(),
    )

    with pytest.raises(ApiError) as exc:
        await service.upload_file(
            supervisor_id=DEV_SUPERVISOR_ID,
            original_filename=original_filename,
            content=b"Valid text.",
            mime_type=mime_type,
        )

    assert exc.value.code == "VALIDATION_ERROR"
    assert exc.value.http_status == 422
    assert exc.value.details[0].field == field
    assert storage.uploaded_paths == []


@pytest.mark.asyncio
async def test_upload_markdown_preserves_source_text_for_retrieval(session: AsyncSession) -> None:
    service = FileService(
        session,
        storage=FakeStorage(),
        embeddings=FakeEmbeddings(),
        settings=Settings(),
    )

    result = await service.upload_file(
        supervisor_id=DEV_SUPERVISOR_ID,
        original_filename="guide.md",
        content=b"# Assessment\n\nUse the published rubric.",
        mime_type="text/markdown",
    )

    chunks = list(
        await session.scalars(select(DocumentChunk).where(DocumentChunk.file_id == result.id))
    )
    assert [chunk.content for chunk in chunks] == ["# Assessment\n\nUse the published rubric."]


@pytest.mark.asyncio
async def test_upload_docx_extracts_paragraphs_and_tables_in_document_order(
    session: AsyncSession,
) -> None:
    document = Document()
    document.add_paragraph("Milestone requirements")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Deliverable"
    table.cell(0, 1).text = "Due"
    table.cell(1, 0).text = "Report"
    table.cell(1, 1).text = "Friday"
    buffer = BytesIO()
    document.save(buffer)
    service = FileService(
        session,
        storage=FakeStorage(),
        embeddings=FakeEmbeddings(),
        settings=Settings(),
    )

    result = await service.upload_file(
        supervisor_id=DEV_SUPERVISOR_ID,
        original_filename="requirements.docx",
        content=buffer.getvalue(),
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    chunks = list(
        await session.scalars(select(DocumentChunk).where(DocumentChunk.file_id == result.id))
    )
    assert [chunk.content for chunk in chunks] == [
        "Milestone requirements\nDeliverable\tDue\nReport\tFriday"
    ]


@pytest.mark.asyncio
async def test_upload_pdf_extracts_text_for_retrieval(session: AsyncSession) -> None:
    service = FileService(
        session,
        storage=FakeStorage(),
        embeddings=FakeEmbeddings(),
        settings=Settings(),
    )

    result = await service.upload_file(
        supervisor_id=DEV_SUPERVISOR_ID,
        original_filename="handbook.pdf",
        content=_pdf_with_text("Capstone PDF guidance"),
        mime_type="application/pdf",
    )

    chunks = list(
        await session.scalars(select(DocumentChunk).where(DocumentChunk.file_id == result.id))
    )
    assert [chunk.content for chunk in chunks] == ["Capstone PDF guidance"]


@pytest.mark.asyncio
@pytest.mark.parametrize("mime_type", ["image/jpeg", "image/png", "image/webp"])
async def test_upload_image_describes_content_before_embedding(
    session: AsyncSession, mime_type: str
) -> None:
    descriptions = FakeImageDescriptions("System diagram: Client connects to API over HTTPS.")
    service = FileService(
        session,
        storage=FakeStorage(),
        embeddings=FakeEmbeddings(),
        image_descriptions=descriptions,
        settings=Settings(),
    )

    result = await service.upload_file(
        supervisor_id=DEV_SUPERVISOR_ID,
        original_filename=f"diagram.{mime_type.rsplit('/', 1)[1]}",
        content=b"image bytes",
        mime_type=mime_type,
    )

    chunks = list(
        await session.scalars(select(DocumentChunk).where(DocumentChunk.file_id == result.id))
    )
    assert descriptions.requests == [(b"image bytes", mime_type)]
    assert [chunk.content for chunk in chunks] == [
        "System diagram: Client connects to API over HTTPS."
    ]


@pytest.mark.asyncio
async def test_upload_marks_image_failed_when_vision_model_fails(session: AsyncSession) -> None:
    service = FileService(
        session,
        storage=FakeStorage(),
        embeddings=FakeEmbeddings(),
        image_descriptions=FailingImageDescriptions(),
        settings=Settings(),
    )

    with pytest.raises(ApiError) as exc:
        await service.upload_file(
            supervisor_id=DEV_SUPERVISOR_ID,
            original_filename="diagram.png",
            content=b"image bytes",
            mime_type="image/png",
        )

    file = await session.scalar(select(File))
    assert exc.value.code == "DOCUMENT_PROCESSING_FAILED"
    assert file is not None
    assert file.processing_status == FileProcessingStatus.FAILED


@pytest.mark.asyncio
async def test_upload_rejects_conflicting_extension_and_mime_type(
    session: AsyncSession,
) -> None:
    storage = FakeStorage()
    service = FileService(
        session,
        storage=storage,
        embeddings=FakeEmbeddings(),
        settings=Settings(),
    )

    with pytest.raises(ApiError) as exc:
        await service.upload_file(
            supervisor_id=DEV_SUPERVISOR_ID,
            original_filename="rubric.pdf",
            content=b"not a PDF",
            mime_type="image/png",
        )

    assert exc.value.code == "UNSUPPORTED_FILE_TYPE"
    assert storage.uploaded_paths == []


@pytest.mark.asyncio
async def test_upload_rejects_pdf_without_extractable_text_before_storage(
    session: AsyncSession,
) -> None:
    storage = FakeStorage()
    service = FileService(
        session,
        storage=storage,
        embeddings=FakeEmbeddings(),
        settings=Settings(),
    )

    with pytest.raises(ApiError) as exc:
        await service.upload_file(
            supervisor_id=DEV_SUPERVISOR_ID,
            original_filename="scan.pdf",
            content=(
                b"%PDF-1.4\n1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n"
                b"2 0 obj<</Type/Pages/Count 0/Kids[]>>endobj\n"
                b"trailer<</Root 1 0 R>>\n%%EOF"
            ),
            mime_type="application/pdf",
        )

    assert exc.value.code == "INVALID_FILE_CONTENT"
    assert storage.uploaded_paths == []
